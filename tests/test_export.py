import io

import openpyxl

from app import export, services

NASTY = ["=HYPERLINK(\"http://evil\",\"x\")", "+cmd|' /C calc'!A0", "-1+1", "@SUM(1+1)", "=1+1",
         "\t=2+2", "http://example.com", "0123", "1e5"]


def test_export_has_zero_formula_cells_and_keeps_text_literal(conn, dishes):
    for i, payload in enumerate(NASTY):
        services.create_dish(conn, title=f"{payload} {i}", description=payload, region=payload, source="user",
                             submitter_visitor=payload, submitter_ip_hash="abc")
    conn.execute("INSERT INTO banned_words(word, word_norm, mode, created_at) VALUES ('=EVIL()', 'evil', 'word', 0)")
    conn.execute("UPDATE categories SET name = '=CMD()' WHERE slug = 'pain'")
    conn.execute("INSERT INTO votes(dish_id, visitor_id, category_id, ip_hash, created_at) VALUES (?,?,1,?,0)",
                 (dishes[0], "=1+2", "+evil"))

    data = export.build_export(conn)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert set(wb.sheetnames) == {"dishes", "votes", "categories", "dish_stats", "submissions", "banned_words"}

    formula_cells = 0
    seen_text = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formula_cells += 1
                if isinstance(cell.value, str):
                    seen_text.add(cell.value)
    assert formula_cells == 0
    for payload in NASTY:
        assert payload in seen_text, payload  # stayed literal text, not converted
    assert "=CMD()" in seen_text and "=EVIL()" in seen_text and "=1+2" in seen_text


def test_export_endpoint_requires_admin(client):
    r = client.get("/admin/export.xlsx", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/admin/login"
